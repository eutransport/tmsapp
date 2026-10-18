"""Tolheffing per land: filteren, opsommen en uitdraaien.

Het land van een tolregel komt uit de GPS-rit, uit een herkende
grensovergang of uit de stationsnaam op de Radius-factuur. Staat het land
nog niet vast, dan blijft het veld leeg. Zulke regels tellen hier bewust
niet mee als buitenland en ook niet als Nederland; ze zijn apart op te
vragen zodat ze niet stilzwijgend bij de verkeerde groep belanden.

Deze module bevat alleen rekenwerk en geen HTTP, zodat de uitdraai en het
overzicht dezelfde selectie gebruiken.
"""
from __future__ import annotations

from decimal import Decimal

from .land_bepaling import LANDNAMEN
from .models import TollingEvent
from .services import build_vehicle_lookup

NEDERLAND = 'NL'
# Twee verzamelkeuzes naast de losse landcodes.
BUITENLAND = 'buitenland'
ONBEKEND = 'onbekend'


def landnaam(code: str) -> str:
    return LANDNAMEN.get(code, code or 'Onbekend')


def filter_op_land(queryset, keuze: str | None):
    """Beperk tolregels tot het gevraagde land.

    Leeg of ``None`` laat alles staan. Verder een ISO-landcode, ``buitenland``
    voor alles buiten Nederland waarvan het land vaststaat, of ``onbekend``
    voor regels zonder land.
    """
    keuze = (keuze or '').strip()
    if not keuze:
        return queryset
    sleutel = keuze.lower()
    if sleutel == BUITENLAND:
        return queryset.exclude(land='').exclude(land=NEDERLAND)
    if sleutel == ONBEKEND:
        return queryset.filter(land='')
    code = keuze.upper()
    if code not in LANDNAMEN:
        toegestaan = ', '.join(sorted(LANDNAMEN) + [BUITENLAND, ONBEKEND])
        raise ValueError(f'Onbekend land "{keuze}". Kies uit: {toegestaan}.')
    return queryset.filter(land=code)


def _plaatsnaam(event: TollingEvent, vloot) -> str:
    """Toon het kenteken zoals het in de vloot staat; anders zoals ingelezen."""
    wagen = vloot.get(event.license_plate_normalized)
    if wagen is not None and wagen.kenteken:
        return wagen.kenteken
    return (event.license_plate_raw or event.license_plate_normalized or '').strip()


def overzicht(queryset) -> list[dict]:
    """Eén regel per kenteken, land en bedrijf, met de ritnummers erbij.

    Het bedrijf staat erbij omdat de uitdraai bedoeld is om te bepalen wie
    de buitenlandse tol doorbelast krijgt.
    """
    events = (queryset
              .select_related('bedrijf')
              .order_by('license_plate_normalized', 'land', 'start_at'))
    vloot = build_vehicle_lookup()

    groepen: dict[tuple, dict] = {}
    for event in events:
        sleutel = (event.license_plate_normalized, event.land, event.bedrijf_id)
        rij = groepen.get(sleutel)
        if rij is None:
            rij = groepen[sleutel] = {
                'plate_normalized': event.license_plate_normalized,
                'plate_display': _plaatsnaam(event, vloot),
                'land': event.land,
                'land_naam': landnaam(event.land),
                'bedrijf_id': str(event.bedrijf_id) if event.bedrijf_id else None,
                'bedrijf_naam': event.bedrijf.naam if event.bedrijf_id else None,
                'ritnummers': [],
                'regels': 0,
                'km': Decimal('0'),
                'bedrag': Decimal('0'),
                'gefactureerd': 0,
                'open_bedrag': Decimal('0'),
                'eerste_dag': None,
                'laatste_dag': None,
            }
        rij['regels'] += 1
        rij['km'] += event.distance_km or Decimal('0')
        rij['bedrag'] += event.amount or Decimal('0')
        if event.invoiced_at is not None:
            rij['gefactureerd'] += 1
        else:
            rij['open_bedrag'] += event.amount or Decimal('0')
        rit = (event.ritnummer or '').strip()
        if rit and rit not in rij['ritnummers']:
            rij['ritnummers'].append(rit)
        dag = event.start_at.date()
        if rij['eerste_dag'] is None or dag < rij['eerste_dag']:
            rij['eerste_dag'] = dag
        if rij['laatste_dag'] is None or dag > rij['laatste_dag']:
            rij['laatste_dag'] = dag

    rijen = list(groepen.values())
    rijen.sort(key=lambda r: (r['plate_display'], r['land'], r['bedrijf_naam'] or ''))
    return rijen


def totalen_per_land(rijen: list[dict]) -> list[dict]:
    per: dict[str, dict] = {}
    for rij in rijen:
        samen = per.setdefault(rij['land'], {
            'land': rij['land'],
            'land_naam': rij['land_naam'],
            'kentekens': set(),
            'regels': 0,
            'km': Decimal('0'),
            'bedrag': Decimal('0'),
            'open_bedrag': Decimal('0'),
        })
        samen['kentekens'].add(rij['plate_normalized'])
        samen['regels'] += rij['regels']
        samen['km'] += rij['km']
        samen['bedrag'] += rij['bedrag']
        samen['open_bedrag'] += rij['open_bedrag']
    uitkomst = []
    for samen in per.values():
        samen['kentekens'] = len(samen['kentekens'])
        uitkomst.append(samen)
    uitkomst.sort(key=lambda r: r['land_naam'])
    return uitkomst


KOPPEN = ['Kenteken', 'Land', 'Bedrijf', 'Ritnummers', 'Eerste dag', 'Laatste dag',
          'Regels', 'Afstand (km)', 'Bedrag (EUR)', 'Gefactureerd', 'Nog open (EUR)']


def naar_xlsx(rijen: list[dict], titel: str) -> bytes:
    """Uitdraai per kenteken met een subtotaalregel per land."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tolheffing buitenland'[:31]

    ws.append([titel])
    ws['A1'].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(KOPPEN)
    kopregel = ws.max_row
    kop_vul = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    for kolom in range(1, len(KOPPEN) + 1):
        cel = ws.cell(row=kopregel, column=kolom)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = kop_vul
        cel.alignment = Alignment(horizontal='center')

    land_vul = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
    huidig_land = object()
    buffer = {'regels': 0, 'km': Decimal('0'), 'bedrag': Decimal('0'),
              'open': Decimal('0'), 'naam': ''}

    def subtotaal():
        if not buffer['regels']:
            return
        ws.append(['', f"Subtotaal {buffer['naam']}", '', '', '', '',
                   buffer['regels'], float(buffer['km']), float(buffer['bedrag']),
                   '', float(buffer['open'])])
        rij = ws.max_row
        for kolom in range(1, len(KOPPEN) + 1):
            cel = ws.cell(row=rij, column=kolom)
            cel.fill = land_vul
            cel.font = Font(bold=True)

    # Per land gesorteerd zodat het subtotaal klopt.
    for rij in sorted(rijen, key=lambda r: (r['land_naam'], r['plate_display'])):
        if rij['land'] != huidig_land:
            subtotaal()
            huidig_land = rij['land']
            buffer = {'regels': 0, 'km': Decimal('0'), 'bedrag': Decimal('0'),
                      'open': Decimal('0'), 'naam': rij['land_naam']}
        ws.append([
            rij['plate_display'],
            rij['land_naam'],
            rij['bedrijf_naam'] or 'Geen bedrijf',
            ', '.join(rij['ritnummers']) or '-',
            rij['eerste_dag'].strftime('%d-%m-%Y') if rij['eerste_dag'] else '',
            rij['laatste_dag'].strftime('%d-%m-%Y') if rij['laatste_dag'] else '',
            rij['regels'],
            float(rij['km']),
            float(rij['bedrag']),
            rij['gefactureerd'],
            float(rij['open_bedrag']),
        ])
        buffer['regels'] += rij['regels']
        buffer['km'] += rij['km']
        buffer['bedrag'] += rij['bedrag']
        buffer['open'] += rij['open_bedrag']
    subtotaal()

    breedtes = [14, 14, 32, 26, 12, 12, 8, 13, 13, 13, 14]
    for kolom, breedte in enumerate(breedtes, start=1):
        ws.column_dimensions[ws.cell(row=kopregel, column=kolom).column_letter].width = breedte

    stroom = io.BytesIO()
    wb.save(stroom)
    return stroom.getvalue()
